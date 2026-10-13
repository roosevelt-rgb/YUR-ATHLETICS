"""Pure Office worker acceptance; every optional reader is test-only.

Run with the shipped interpreter for mandatory extractor/OPC/security checks.
Running this same suite with a development interpreter containing python-docx,
openpyxl and python-pptx additionally proves independent reader interoperability;
none of those packages is a product runtime dependency or a native Office claim.
"""
import copy
import hashlib
import importlib.util
import io
import json
import os
from pathlib import Path
import posixpath
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
from xml.etree import ElementTree as ET
from zipfile import ZIP_STORED, ZipFile

from foundry.attachments import read_attachment
from foundry.specialists import validate_schema
from foundry.work_documents import artifact_bytes, catalog, create_document


REPORT = {"name": "project-brief.docx", "format": "docx", "title": "Project <A> & B",
          "sections": [{"heading": "Scope", "paragraphs": ["125 CY of concrete.", "Cost: $18,750.50", "Line one\nLine two"]},
                       {"heading": "Decision", "paragraphs": ["Proceed only after source review.", "<script>alert('inert')</script>"]}]}
WORKBOOK = {"name": "takeoff.xlsx", "format": "xlsx", "title": "Quantity review",
            "sheets": [{"name": "Takeoff", "columns": ["Item", "Quantity", "Approved", "Comment"],
                        "rows": [["Concrete", 125.5, True, None], ["No allowance", 0, False, "=SUM(B2:B3)"],
                                 ["<xml> & café", -25, None, "https://example.invalid/data"]]},
                       {"name": "Rates", "columns": ["Unit", "Rate"], "rows": [["CY", 150]]}]}
DECK = {"name": "delivery-review.pptx", "format": "pptx", "title": "Delivery review",
        "slides": [{"title": "Scope & evidence", "bullets": ["125 CY confirmed", "Cost: $18,750.50", "<script>literal data</script>"]},
                   {"title": "Decision", "bullets": ["Review the source before applying edits."]}]}


class WorkDocumentTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def produce(self, recipe):
        result = create_document(copy.deepcopy(recipe))
        raw = artifact_bytes(result["proposal"])
        self.assertEqual(hashlib.sha256(raw).hexdigest(), result["proposal"]["sha256"])
        self.assertEqual(len(raw), result["proposal"]["bytes"])
        return result["proposal"], raw

    def acquire(self, recipe):
        proposal, raw = self.produce(recipe)
        path = self.root / proposal["name"]
        path.write_bytes(raw)
        return read_attachment({"path": str(path), "expected_sha256": proposal["sha256"]})

    def test_typed_catalog_and_compact_recipe_identity(self):
        descriptor = catalog()[0]
        self.assertEqual(descriptor["name"], "create_document")
        self.assertEqual(descriptor["permissions"], [])
        self.assertEqual(descriptor["runtime"]["dependencies"], [])
        self.assertFalse(descriptor["runtime"]["network"])
        self.assertEqual(len(descriptor["input_schema"]["oneOf"]), 3)
        for recipe in (REPORT, WORKBOOK, DECK):
            with self.subTest(format=recipe["format"]):
                validate_schema(recipe, descriptor["input_schema"])
                result = create_document(recipe)
                validate_schema(result, descriptor["output_schema"])
                proposal = result["proposal"]
                self.assertEqual(json.loads(proposal["content"]), recipe)
                self.assertEqual(proposal["encoding"], "ooxml-spec-v1")
                self.assertLess(len(proposal["content"]), proposal["bytes"])
                self.assertNotIn("base64", proposal)
        descriptor["input_schema"]["properties"].clear()
        self.assertIn("format", catalog()[0]["input_schema"]["properties"])

    def test_existing_attachment_intake_reads_exact_report(self):
        result = self.acquire(REPORT)
        self.assertEqual(result["content"], "Project <A> & B\nScope\n125 CY of concrete.\nCost: $18,750.50\nLine one\nLine two\nDecision\nProceed only after source review.\n<script>alert('inert')</script>")
        self.assertFalse(result["provenance"]["code_executed"])
        self.assertFalse(result["extraction"]["truncated"])

    def test_existing_attachment_intake_preserves_numeric_cells_and_formula_text(self):
        result = self.acquire(WORKBOOK)
        self.assertIn("A2=Concrete\tB2=125.5\tC2=1\tD2=", result["content"])
        self.assertIn("A3=No allowance\tB3=0\tC3=0\tD3==SUM(B2:B3)", result["content"])
        self.assertIn("A4=<xml> & café\tB4=-25\tC4=\tD4=https://example.invalid/data", result["content"])
        self.assertIn("A2=CY\tB2=150", result["content"])
        self.assertFalse(result["extraction"]["formulas_executed"])
        self.assertNotIn("[formula:", result["content"])
        self.assertEqual(result["extraction"]["sheets"], 2)

    def test_existing_attachment_intake_preserves_all_slides(self):
        result = self.acquire(DECK)
        self.assertEqual(result["content"], "[ppt/slides/slide1.xml]\nScope & evidence\n125 CY confirmed\nCost: $18,750.50\n<script>literal data</script>\n\n[ppt/slides/slide2.xml]\nDecision\nReview the source before applying edits.")
        self.assertEqual(result["extraction"]["slides"], 2)

    def test_opc_parts_relationships_content_types_and_fixed_metadata(self):
        rel_ns = "{http://schemas.openxmlformats.org/package/2006/relationships}"
        for recipe in (REPORT, WORKBOOK, DECK):
            with self.subTest(format=recipe["format"]):
                _, raw = self.produce(recipe)
                with ZipFile(io.BytesIO(raw)) as package:
                    names = package.namelist()
                    self.assertEqual(names, sorted(names))
                    self.assertEqual(len(names), len(set(names)))
                    self.assertTrue(all(item.date_time == (2000, 1, 1, 0, 0, 0) for item in package.infolist()))
                    self.assertTrue(all(item.compress_type == ZIP_STORED for item in package.infolist()))
                    types = ET.fromstring(package.read("[Content_Types].xml"))
                    overrides = {item.get("PartName").lstrip("/") for item in types if item.tag.endswith("Override")}
                    for name in names:
                        xml = ET.fromstring(package.read(name))
                        if not name.endswith(".rels") and name != "[Content_Types].xml":
                            self.assertIn(name, overrides)
                        if name.endswith(".rels"):
                            base = "" if name == "_rels/.rels" else posixpath.dirname(posixpath.dirname(name))
                            ids = []
                            for rel in xml.findall(rel_ns + "Relationship"):
                                ids.append(rel.get("Id"))
                                self.assertIsNone(rel.get("TargetMode"))
                                self.assertNotIn(":", rel.get("Target"))
                                self.assertIn(posixpath.normpath(posixpath.join(base, rel.get("Target"))), names)
                            self.assertEqual(len(ids), len(set(ids)))
                    core = package.read("docProps/core.xml").decode()
                    self.assertIn("YUR Desktop", core)
                    self.assertEqual(core.count("2000-01-01T00:00:00Z"), 2)
                    self.assertNotIn(str(self.root), core)

    def test_xml_injection_stays_text_and_has_no_executable_parts(self):
        forbidden_tags = {"script", "f", "oleObject", "hyperlink", "hlinkClick", "externalLink", "altChunk"}
        for recipe in (REPORT, WORKBOOK, DECK):
            _, raw = self.produce(recipe)
            with ZipFile(io.BytesIO(raw)) as package:
                self.assertFalse(any("vba" in name.lower() or name.endswith(".bin") for name in package.namelist()))
                for name in package.namelist():
                    text = package.read(name)
                    self.assertNotIn(b"<!DOCTYPE", text)
                    for node in ET.fromstring(text).iter():
                        self.assertNotIn(node.tag.rsplit("}", 1)[-1], forbidden_tags)
                if recipe["format"] == "xlsx":
                    self.assertIn(b'<c r="D3" s="0" t="inlineStr">', package.read("xl/worksheets/sheet1.xml"))

    def test_determinism_reordered_input_and_fresh_process(self):
        for recipe in (REPORT, WORKBOOK, DECK):
            proposal, raw = self.produce(recipe)
            reversed_keys = dict(reversed(list(recipe.items())))
            self.assertEqual(create_document(reversed_keys)["proposal"], proposal)
            command = "import json,sys; from foundry.work_documents import create_document; print(create_document(json.load(sys.stdin))['proposal']['sha256'])"
            for seed in ("7", "23"):
                process = subprocess.run([sys.executable, "-B", "-c", command], input=json.dumps(recipe), text=True,
                                         capture_output=True, cwd=Path(__file__).resolve().parents[1],
                                         env={**os.environ, "PYTHONHASHSEED": seed, "PYTHONNOUSERSITE": "1"}, timeout=20, check=True)
                self.assertEqual(process.stdout.strip(), proposal["sha256"])
                self.assertEqual(process.stderr, "")
            self.assertEqual(artifact_bytes(proposal), raw)

    def test_worker_performs_no_filesystem_network_or_process_operations(self):
        original = copy.deepcopy(WORKBOOK)
        with patch("builtins.open", side_effect=AssertionError("unexpected file operation")), \
             patch("socket.socket", side_effect=AssertionError("unexpected network operation")), \
             patch("subprocess.Popen", side_effect=AssertionError("unexpected process operation")):
            proposal = create_document(WORKBOOK)["proposal"]
            self.assertGreater(len(artifact_bytes(proposal)), 100)
        self.assertEqual(WORKBOOK, original)

    def test_exact_format_body_and_unknown_arguments_fail_closed(self):
        cases = []
        for recipe in (REPORT, WORKBOOK, DECK):
            body = {"docx": "sections", "xlsx": "sheets", "pptx": "slides"}[recipe["format"]]
            missing = copy.deepcopy(recipe)
            del missing[body]
            cases.extend([missing, {**recipe, "root": "C:/unrelated"}, {**recipe, "authority": {"commit": True}}, {**recipe, "template": "https://example.invalid"}])
        cases.append({**REPORT, "slides": DECK["slides"]})
        cases.append({**REPORT, "format": "pdf"})
        nested = copy.deepcopy(DECK)
        nested["slides"][0]["script"] = "execute"
        cases.append(nested)
        for recipe in cases:
            with self.subTest(recipe=recipe), self.assertRaises(ValueError):
                create_document(recipe)

    def test_malicious_and_mismatched_names_rejected(self):
        for name in ("../report.docx", "C:\\report.docx", "folder/report.docx", "CON.docx", "LPT1.docx", ".hidden.docx",
                     "report.docx:stream", "report.xlsx", "credentials.docx", "report.docx ", "report?.docx", "a\x00.docx", "a" * 121 + ".docx"):
            with self.subTest(name=name), self.assertRaises(ValueError):
                create_document({**REPORT, "name": name})

    def test_secret_control_and_surrogate_content_rejected(self):
        for value in ("-----BEGIN PRIVATE KEY-----", "password = 'fixture-credential-value'", "Bearer abcdefghijklmnop12345",
                      "https://username:passphrase@example.invalid", "bad\x00text", "bad\rtext", "bad\ttext", "bad\ud800", "bad\ufffe"):
            recipe = copy.deepcopy(REPORT)
            recipe["sections"][0]["paragraphs"] = [value]
            with self.subTest(value=repr(value)), self.assertRaises(ValueError):
                create_document(recipe)

    def test_shape_and_size_limits_never_silently_truncate_data(self):
        oversized = copy.deepcopy(REPORT)
        oversized["sections"][0]["paragraphs"] = ["x" * 4001]
        total = copy.deepcopy(REPORT)
        total["sections"] = [{"heading": "Large", "paragraphs": ["é" * 3500] * 10}]
        slides = {**DECK, "slides": DECK["slides"] * 7}
        rows = {**WORKBOOK, "sheets": [{"name": "Large", "columns": [f"C{i}" for i in range(20)], "rows": [[0] * 20] * 100}]}
        for recipe in (oversized, total, slides, rows):
            with self.assertRaises(ValueError):
                create_document(recipe)
        with patch("foundry.work_documents.MAX_ARTIFACT_BYTES", 20), self.assertRaises(ValueError):
            create_document(REPORT)
        preview = {**REPORT, "sections": [{"heading": "Long", "paragraphs": ["x" * 3000, "y" * 3000]}]}
        proposal, _ = self.produce(preview)
        self.assertEqual(len(proposal["preview_text"]), 4000)
        self.assertTrue(proposal["preview_truncated"])
        self.assertEqual(len(json.loads(proposal["content"])["sections"][0]["paragraphs"][1]), 3000)

    def test_spreadsheet_shape_names_numbers_and_literal_formula_values(self):
        for value in (float("nan"), float("inf"), -float("inf"), 10 ** 13, 1e-310, 0.1234567890123456, {"formula": "A1+1"}):
            recipe = copy.deepcopy(WORKBOOK)
            recipe["sheets"][0]["rows"][0][1] = value
            with self.subTest(value=value), self.assertRaises(ValueError):
                create_document(recipe)
        for name in ("Bad/Name", "Bad[Name", "'quoted", "History", "Takeoff", " "):
            recipe = copy.deepcopy(WORKBOOK)
            recipe["sheets"][1]["name"] = name
            with self.subTest(name=name), self.assertRaises(ValueError):
                create_document(recipe)
        recipe = copy.deepcopy(WORKBOOK)
        recipe["sheets"][0]["rows"][0].pop()
        with self.assertRaises(ValueError):
            create_document(recipe)
        recipe = copy.deepcopy(WORKBOOK)
        recipe["sheets"][0]["columns"] = ["Name", "name", "Other", "End"]
        with self.assertRaises(ValueError):
            create_document(recipe)

    def test_artifact_regeneration_rejects_any_changed_identity_or_noncanonical_recipe(self):
        proposal, _ = self.produce(REPORT)
        mutations = {"sha256": "f" * 64, "bytes": proposal["bytes"] + 1, "name": "different.docx",
                     "media_type": "text/plain", "encoding": "base64", "preview_text": "forged preview",
                     "preview_truncated": True, "content": json.dumps(REPORT)}
        for key, value in mutations.items():
            with self.subTest(key=key), self.assertRaises(ValueError):
                artifact_bytes({**proposal, key: value})
        with self.assertRaises(ValueError):
            artifact_bytes({**proposal, "path": "C:/unrelated"})
        altered = json.loads(proposal["content"])
        altered["title"] = "Different actual document"
        altered = json.dumps(altered, sort_keys=True, ensure_ascii=False, separators=(",", ":"))
        with self.assertRaises(ValueError):
            artifact_bytes({**proposal, "content": altered})
        for content in ('{"name":"a.docx","name":"b.docx"}', "{" + '"a":' + "[" * 2000 + "0" + "]" * 2000 + "}"):
            with self.assertRaises(ValueError):
                artifact_bytes({**proposal, "content": content})

    @unittest.skipUnless(importlib.util.find_spec("docx"), "Test-only python-docx is not installed in this interpreter")
    def test_independent_docx_reader(self):
        from docx import Document
        _, raw = self.produce(REPORT)
        document = Document(io.BytesIO(raw))
        self.assertEqual(document.core_properties.title, REPORT["title"])
        self.assertEqual([paragraph.text for paragraph in document.paragraphs], self.acquire(REPORT)["content"].split("\n"))
        self.assertEqual(document.paragraphs[0].style.name, "Title")
        self.assertEqual(document.paragraphs[1].style.style_id, "Heading1")

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "Test-only openpyxl is not installed in this interpreter")
    def test_independent_xlsx_reader(self):
        from openpyxl import load_workbook
        _, raw = self.produce(WORKBOOK)
        workbook = load_workbook(io.BytesIO(raw), data_only=False)
        self.assertEqual(workbook.properties.title, WORKBOOK["title"])
        self.assertEqual(workbook.sheetnames, ["Takeoff", "Rates"])
        for expected in WORKBOOK["sheets"]:
            sheet = workbook[expected["name"]]
            self.assertEqual(list(sheet.values), [tuple(expected["columns"]), *(tuple(row) for row in expected["rows"])])
            self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(workbook["Takeoff"]["D3"].data_type, "s")
        self.assertTrue(workbook["Takeoff"]["A1"].font.bold)
        workbook.close()

    @unittest.skipUnless(importlib.util.find_spec("pptx"), "Test-only python-pptx is not installed in this interpreter")
    def test_independent_pptx_reader(self):
        from pptx import Presentation
        _, raw = self.produce(DECK)
        deck = Presentation(io.BytesIO(raw))
        self.assertEqual(deck.core_properties.title, DECK["title"])
        self.assertEqual(len(deck.slides), 2)
        self.assertEqual(len(deck.slide_masters), 1)
        self.assertEqual(len(deck.slide_layouts), 1)
        for slide, expected in zip(deck.slides, DECK["slides"]):
            self.assertEqual(slide.shapes.title.text, expected["title"])
            self.assertEqual(slide.shapes[1].text, "\n".join(expected["bullets"]))


if __name__ == "__main__":
    unittest.main()
